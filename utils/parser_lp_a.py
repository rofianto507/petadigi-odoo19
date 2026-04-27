import re
import base64
import io
from datetime import datetime

def parse(file_bytes: bytes) -> dict:
    """
    Parse dokumen LP A (.docx) dan kembalikan dict berisi field-field
    yang berhasil diekstrak.
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "Library python-docx tidak ditemukan. "
            "Silakan install dengan: pip install python-docx"
        )

    doc = Document(io.BytesIO(file_bytes))
    result = {}

    # Kumpulkan semua teks paragraf
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(paragraphs)

    # ── NO. LP ──────────────────────────────────────────────────────────────
    # Pola: "Nomor :" atau "Nomor:" diikuti teks LP
    for para in paragraphs:
        m = re.match(r'Nomor\s*:\s*(.+)', para, re.IGNORECASE)
        if m:
            result['no_lp'] = m.group(1).strip()
            break

    # ── PERISTIWA (numbered list) ────────────────────────────────────────────
    # 1. Waktu Kejadian
    m = re.search(
        r'1\.?\s*Waktu Kejadian\s*:?\s*(.+?)(?=\n2\.)',
        full_text, re.IGNORECASE | re.DOTALL
    )
    if m:
        result['tanggal_kejadian_raw'] = m.group(1).strip()
        result['tanggal_kejadian'] = _parse_tanggal(m.group(1).strip())

    # 2. Tempat Kejadian (termasuk koordinat)
    m = re.search(
        r'2\.?\s*Tempat Kejadian\s*:?\s*(.+?)(?=\n3\.)',
        full_text, re.IGNORECASE | re.DOTALL
    )
    if m:
        tempat = m.group(1).strip()
        result['tempat_kejadian'] = tempat

        # Ekstrak TITIK KOORDINAT lat,lon
        coord = re.search(
            r'TITIK KOORDINAT\s*([\-\d\.]+)[\s,]+([\-\d\.]+)',
            tempat, re.IGNORECASE
        )
        if coord:
            try:
                result['latitude'] = float(coord.group(1))
                result['longitude'] = float(coord.group(2))
            except ValueError:
                pass

    # 3. Apa Yang Terjadi
    m = re.search(
        r'3\.?\s*Apa\s+[Yy]ang\s+[Tt]erjadi\s*:?\s*(.+?)(?=\n4\.)',
        full_text, re.IGNORECASE | re.DOTALL
    )
    if m:
        result['apa_yang_terjadi'] = m.group(1).strip()

    # 4. Terlapor & Korban
    m = re.search(
        r'4\.?\s*Siapa\s*:?(.+?)(?=\n5\.)',
        full_text, re.IGNORECASE | re.DOTALL
    )
    if m:
        blok = m.group(1)
        # Terlapor
        t = re.search(r'Terlapor\s*:?\s*(.+?)(?=Korban\s*:|$)',
                      blok, re.IGNORECASE | re.DOTALL)
        if t:
            result['terlapor'] = t.group(1).strip()
        # Korban
        k = re.search(r'Korban\s*:?\s*(.+?)$',
                      blok, re.IGNORECASE | re.DOTALL)
        if k:
            val = k.group(1).strip()
            if val and val != '-':
                result['korban'] = val

    # 5. Bagaimana Terjadi -> uraian_singkat
    m = re.search(
        r'5\.?\s*Bagaimana\s+[Tt]erjadi\s*:?\s*(.+?)(?=\n6\.)',
        full_text, re.IGNORECASE | re.DOTALL
    )
    if m:
        result['uraian_singkat'] = m.group(1).strip()

    # 6. Dilaporkan Pada -> tanggal_laporan
    m = re.search(
        r'6\.?\s*Dilaporkan\s+Pada\s*:?\s*(.+?)(?=\n|$)',
        full_text, re.IGNORECASE | re.DOTALL
    )
    if m:
        result['tanggal_laporan_raw'] = m.group(1).strip()
        result['tanggal_laporan'] = _parse_tanggal(m.group(1).strip())

    # ── TABEL ────────────────────────────────────────────────────────────────
    # Struktur tabel LP A:
    # Baris 1: [TINDAK PIDANA APA | NAMA DAN ALAMAT SAKSI-SAKSI]
    # Baris 2: [isi tindak pidana | isi saksi]
    # Baris 3: [BARANG BUKTI | URAIAN SINGKAT YANG DILAPORKAN]
    # Baris 4: [isi barang bukti | isi uraian singkat]
    for table in doc.tables:
        teks_sel = [[cell.text.strip() for cell in row.cells] for row in table.rows]

        for i, row in enumerate(teks_sel):
            # Header TINDAK PIDANA
            if any('TINDAK PIDANA' in c.upper() for c in row):
                if i + 1 < len(teks_sel):
                    next_row = teks_sel[i + 1]
                    if len(next_row) >= 1:
                        result['tindak_pidana'] = next_row[0].strip()
                    if len(next_row) >= 2:
                        result['saksi'] = next_row[1].strip()

            # Header BARANG BUKTI
            if any('BARANG BUKTI' in c.upper() for c in row):
                if i + 1 < len(teks_sel):
                    next_row = teks_sel[i + 1]
                    if len(next_row) >= 1:
                        result['barang_bukti'] = next_row[0].strip()
                    if len(next_row) >= 2:
                        # Ambil uraian singkat dari tabel jika belum ada
                        uraian_tabel = next_row[1].strip()
                        if uraian_tabel and 'uraian_singkat' not in result:
                            result['uraian_singkat'] = uraian_tabel

    # ── PELAPOR (dari footer / tanda tangan) ────────────────────────────────
    # Cari pola: paragraf "Pelapor" diikuti nama
    for idx, para in enumerate(paragraphs):
        if re.match(r'^Pelapor\s*$', para, re.IGNORECASE):
            # Ambil 2 baris setelah "Pelapor"
            nama_lines = []
            for j in range(idx + 1, min(idx + 4, len(paragraphs))):
                line = paragraphs[j].strip()
                if line:
                    nama_lines.append(line)
                if len(nama_lines) == 2:
                    break
            if nama_lines:
                result['pelapor'] = '\n'.join(nama_lines)
            break

    return result


# ── HELPER ──────────────────────────────────────────────────────────────────

BULAN = {
    'januari': 1, 'februari': 2, 'maret': 3, 'april': 4,
    'mei': 5, 'juni': 6, 'juli': 7, 'agustus': 8,
    'september': 9, 'oktober': 10, 'november': 11, 'desember': 12,
}

def _parse_tanggal(teks: str):
    """
    Coba parse teks tanggal Bahasa Indonesia ke datetime.
    Contoh:
      'senin tanggal 16 Februari 2026'
      'Pada Hari Senin Tanggal 16 Februari 2026 Pukul 19.25 WIB'
    Kembalikan datetime atau None jika gagal.
    """
    teks_lower = teks.lower()

    # Ekstrak jam jika ada (pukul HH.MM)
    jam, menit = 0, 0
    m_jam = re.search(r'pukul\s+(\d{1,2})[.:\-](\d{2})', teks_lower)
    if m_jam:
        jam = int(m_jam.group(1))
        menit = int(m_jam.group(2))

    # Cari pola: (tanggal)? <angka> <bulan> <tahun>
    m = re.search(
        r'(\d{1,2})\s+(' + '|'.join(BULAN.keys()) + r')\s+(\d{4})',
        teks_lower
    )
    if m:
        try:
            return datetime(
                int(m.group(3)),
                BULAN[m.group(2)],
                int(m.group(1)),
                jam, menit
            )
        except ValueError:
            pass
    return None
